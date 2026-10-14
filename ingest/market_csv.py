"""v3.3 connector sàn (P1, CHỈ ĐỌC) — nạp file CSV/XLSX Quản lý MXH xuất tay
từ Seller Center Shopee / TikTok Shop vào staging, đối chiếu golden (v3.2)
bằng QUY TẮC (không LLM) và mở/đóng cảnh báo trong `market_alerts`.

Ownership / ranh giới:
  - Ghi DUY NHẤT 3 bảng `market_snapshots`, `market_rows`, `market_alerts`.
  - CHỈ ĐỌC golden (`sku_master`, `price_list`) — không ghi golden, không
    ghi `products`, không gọi sàn, không đổi giá/listing/tồn kho/đơn. Người
    sửa trên sàn là Quản lý MXH sau khi GĐ bán hàng (giá) / Pháp chế (tuân
    thủ) duyệt.

Contract nạp (`python -m ingest.market_csv --file F --channel shopee|tiktok
--kind listing|order|ads`):
  - Parse + kiểm TOÀN file TRƯỚC khi connect (`read_export`, hàm thuần).
    Đuôi không phải .csv/.xlsx, CSV không decode `utf-8-sig`, file rỗng /
    chỉ header, thiếu cột `required`, cột lạ (ngoài required/optional/drop),
    header trùng, số không parse được số nguyên ≥ 0, dòng thiếu cả item_id
    lẫn seller_sku, order thiếu mã đơn, vượt MAX_ROWS dòng / MAX_COLS cột
    (áp trong lúc duyệt) → in lỗi (chỉ số dòng + tên field,
    KHÔNG in giá trị ô), exit 1, KHÔNG connect, 0 dòng nạp.
  - `file_sha256` đã có → `[skip]`, exit 0, không ghi. File mới → 1
    transaction: snapshot + rows + `reconcile()` → `conn.commit()`.

Mapping cột: `data/market_csv_map.json` (tên cột GIẢ ĐỊNH — chưa xác minh với
file thật). `{kênh: {loại: {required: {field: header}, optional: {...},
drop: [header], inactive_status: [...], cancel_status: [...]}}}`. Field
chuẩn: item_id, seller_sku, title, price, stock, status, ads_spend, qty,
order_id, sale_price (v5.3: giá sau khuyến mãi của listing, TÙY CHỌN — chỉ
nằm trong `market_rows.raw`, không có cột riêng; cảnh báo v3.3 vẫn so theo
`price`; agents.kd_price đọc làm giá bán thực), voucher_code (v5.4: mã
voucher/affiliate trên dòng đơn, TÙY CHỌN, ≤ VOUCHER_MAX ký tự — chỉ trong
`raw`; agents.kd_affiliate ghi nhận đơn qua mã). PII: chỉ field allowlist
được giữ; cột `drop` (tên/SĐT/địa chỉ người mua, ghi chú...) bị bỏ ngay khi đọc; `order_id` → `order_id_hash` =
sha256("kênh:mã đơn")[:16], giá trị gốc không giữ.

Ánh xạ item sàn → golden: (1) `seller_sku` = `sku_master.sku_key`; (2) kênh
shopee và `item_id` khớp ĐÚNG 1 `sku_master.shopee_item_id`; khớp ≥2 (biến
thể chung item) hoặc 0 → không map (detail ghi ứng viên, agent không chọn).

Cảnh báo (`reconcile`, chỉ listing/ads, dựng lại trên snapshot vừa nạp của
cùng (kênh, loại); cảnh báo open mà item CÓ MẶT trong file mới và hết lỗi →
closed; item VẮNG mặt → giữ open + `[warn]` (vắng ≠ đã sửa)):
  alcohol_listed     S1  listing map tới SKU rượu, status không thuộc
                         inactive_status (status lạ = đang bán, fail-closed)
  price_mismatch     S2  listing giá ≠ giá `list` golden đang mở cùng kênh
                         (detail ghi golden_approved true/false)
  tpbvsk_ads_no_xnqc S1  ads map tới TPBVSK chưa XNQC approved, chi phí > 0
  item_unmapped      S2  dòng listing/ads không map được golden
Export `order` chỉ nạp staging (báo cáo tách đơn hủy), không sinh cảnh báo.

`--report`: cảnh báo open S1 trước S2, snapshot mới nhất mỗi (kênh, loại)
+ tuổi ngày, "[warn] <kênh>: không đủ dữ liệu quảng cáo" khi có listing mà
chưa có file ads, đơn theo trạng thái (đơn hủy riêng). Không in mã đơn/hash.
"""

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import unicodedata
from itertools import islice
from pathlib import Path

import psycopg
from dotenv import load_dotenv
from psycopg.types.json import Jsonb

from ingest.golden_store import TODAY_VN
from pipelines.pii import _mask_pii

# Console Windows mặc định cp1258 — crash UnicodeEncodeError khi in tiếng
# Việt có dấu. Ép UTF-8 cho stdout/stderr (no-op nếu đã là UTF-8).
for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
MAP_FILE = ROOT / "data" / "market_csv_map.json"

CHANNELS = ("shopee", "tiktok")
KINDS = ("listing", "order", "ads")
FIELDS = ("item_id", "seller_sku", "title", "price", "stock", "status",
          "ads_spend", "qty", "order_id", "sale_price", "voucher_code")
# v5.4: mã voucher/affiliate trên dòng đơn (TÙY CHỌN, chỉ vào `raw`) — trần
# độ dài; vượt → lỗi dòng (không in giá trị ô).
VOUCHER_MAX = 64
# field số → trần giá trị (cột bigint tiền VND / int số lượng).
NUM_MAX = {"price": 10**15, "ads_spend": 10**15, "stock": 2 * 10**9,
           "qty": 2 * 10**9, "sale_price": 10**15}
MAX_FILE_BYTES = 50 * 1024 * 1024
# Trần duyệt (áp TRONG lúc đọc, không dựng cả bảng trong RAM): MAX_FILE_BYTES
# chỉ chặn byte trên đĩa — xlsx 5KB có dimension A1:XFD1048576 hay CSV toàn
# dòng trống vẫn nở ra hàng GB nếu list() hết. Dòng đếm cả dòng trống.
MAX_ROWS = 100_000
MAX_COLS = 200
MAX_PRINT_ERRS = 30

# Số nguyên VND: "159000" | nhóm 3 chữ số cùng 1 ký tự phân cách ("482.900",
# "1,490,000"). "12,5" / "1.234,50" (phần thập phân) → lỗi, không đoán.
_INT_RE = re.compile(r"^(?:[0-9]+|[0-9]{1,3}(?:([.,])[0-9]{3})(?:\1[0-9]{3})*)$")
_STRIP_RE = re.compile(r"[\s₫]|(?:VND|vnđ|đ)$", re.IGNORECASE)

SEVERITY = {"alcohol_listed": "S1", "tpbvsk_ads_no_xnqc": "S1",
            "price_mismatch": "S2", "item_unmapped": "S2"}


def _nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s).strip()


def _clean(s, n: int = 120) -> str:
    """Chuỗi từ file sàn trước khi in: bỏ ký tự điều khiển (escape
    terminal), cắt độ dài."""
    return "".join(c if c.isprintable() else "?" for c in str(s))[:n]


def _safe(s, n: int = 60) -> str:
    """Header lạ trước khi in: che SĐT/email rồi `_clean`."""
    return _clean(_mask_pii(str(s)), n)


def load_mapping(path: Path = MAP_FILE) -> tuple[dict | None, str | None]:
    try:
        return json.loads(path.read_text(encoding="utf-8")), None
    except (OSError, ValueError) as e:
        return None, f"không đọc được mapping {path.name}: {type(e).__name__}"


def _cell(v) -> str | None:
    """Ô CSV/XLSX → str đã strip, ô rỗng → None. Số nguyên xlsx (int /
    float nguyên) → chữ số thuần (item id, giá). Float CÓ phần lẻ → dạng
    `f` 6 chữ số thập phân ("12.345000") — KHÔNG dùng repr: repr(12.345) =
    "12.345" khớp nhánh nhóm-nghìn của _INT_RE → 12345 (×1000 câm). Dạng
    "X.dddddd" không bao giờ khớp _INT_RE → lỗi số, fail-closed."""
    if v is None:
        return None
    if isinstance(v, float):
        v = str(int(v)) if v.is_integer() else f"{v:f}"
    else:
        v = str(v)   # int/bool/str/datetime
    v = v.strip()
    return v or None


def _to_int(s: str | None, cap: int) -> int | None:
    if s is None:
        return None
    s = _STRIP_RE.sub("", s)
    # chặn chuỗi số dài trước int() (Python giới hạn 4300 chữ số → ValueError)
    if len(s) > 25 or not _INT_RE.match(s):
        return None
    n = int(s.replace(".", "").replace(",", ""))
    return n if n <= cap else None


class _TableError(Exception):
    """Lỗi cấu trúc file — message không chứa nội dung ô."""


def _raw_rows(data: bytes, suffix: str):
    """Duyệt LƯỜI từng dòng thô (CSV list / XLSX tuple, tối đa MAX_COLS+1
    cột). Lỗi đọc → _TableError."""
    if suffix == ".csv":
        try:
            data.decode("utf-8-sig")   # kiểm strict, bỏ kết quả ngay
        except UnicodeDecodeError as e:
            raise _TableError(f"file CSV không phải UTF-8 (byte {e.start})"
                              ) from None
        # decode lười theo chunk — StringIO(text) giữ thêm bản UCS-4 (4×)
        text = io.TextIOWrapper(io.BytesIO(data), encoding="utf-8-sig",
                                newline="")
        try:
            yield from csv.reader(text)
        except csv.Error as e:
            # thông báo csv.Error không chứa nội dung ô (vd field limit)
            raise _TableError(f"CSV hỏng: {e}"[:160]) from None
        return
    try:
        import openpyxl  # lười: chỉ cần khi nạp .xlsx
    except ImportError:
        raise _TableError("thiếu openpyxl — cài `pip install openpyxl` hoặc "
                          "xuất CSV") from None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                    data_only=True)
    except Exception as e:  # noqa: BLE001 — zip/xml hỏng ném nhiều loại;
        # fail-closed: báo lỗi, 0 dòng nạp, chỉ nêu loại lỗi không nội dung
        raise _TableError(f"không đọc được XLSX: {type(e).__name__}") from None
    try:
        ws = wb.worksheets[0]
        # used-range (dimension) phóng đại → read-only openpyxl đệm dòng
        # rỗng cho mọi khoảng trống → chặn TRƯỚC khi duyệt.
        if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS:
            raise _TableError(
                f"vùng dữ liệu sheet {ws.max_row} dòng × {ws.max_column} cột"
                f" vượt trần {MAX_ROWS} × {MAX_COLS} — xóa dòng/cột thừa "
                "(ô định dạng xa phía dưới/bên phải) hoặc xuất CSV")
        # dimension KHAI THIẾU (A1:F2 mà có dòng 4) → read-only dừng ở
        # max_row = cắt câm dòng cuối (có thể là dòng rượu → mất S1).
        # Bỏ dimension, đọc tới hết sheet; trần MAX_ROWS/MAX_COLS vẫn áp
        # trong lúc duyệt ở `_table` (dòng đệm dùng chung 1 tuple).
        ws.reset_dimensions()
        yield from ws.iter_rows(values_only=True, max_col=MAX_COLS + 1)
    except _TableError:
        raise
    except Exception as e:  # noqa: BLE001 — read-only parse lười, lỗi xml
        raise _TableError(f"không đọc được XLSX: {type(e).__name__}") from None
    finally:
        wb.close()


def _table(data: bytes, suffix: str):
    """→ iterator (số dòng file, [ô đã `_cell`, bỏ None cuối]) — bỏ dòng
    trống NGAY, áp trần MAX_ROWS (đếm cả dòng trống) / MAX_COLS trong lúc
    duyệt. Không bao giờ giữ cả bảng thô trong RAM."""
    for i, r in enumerate(_raw_rows(data, suffix), 1):
        if i > MAX_ROWS:
            raise _TableError(f"file > {MAX_ROWS} dòng — tách file")
        if len(r) > MAX_COLS and any(c not in (None, "")
                                     for c in islice(r, MAX_COLS, None)):
            raise _TableError(f"dòng {i}: có ô ngoài {MAX_COLS} cột")
        cells = [_cell(c) for c in islice(r, MAX_COLS)]
        while cells and cells[-1] is None:
            cells.pop()
        if cells:
            yield i, cells


def read_export(path: Path, channel: str, kind: str, mapping: dict
                ) -> tuple[list[dict], list[str], str | None]:
    """Hàm thuần (không DB): → (rows, errors, sha256 bytes file).
    errors ≠ [] → rows = [] (không nạp nửa vời). Mỗi row CHỈ gồm field
    allowlist + `row_no` (số dòng trong file, header = 1); `order_id` đã
    đổi thành `order_id_hash`. Lỗi chỉ nêu số dòng/tên cột, không giá trị ô.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        return [], [f"đuôi '{suffix}' không hỗ trợ — chỉ .csv / .xlsx"], None
    spec = (mapping.get(channel) or {}).get(kind)
    if not isinstance(spec, dict) or not isinstance(spec.get("required"), dict):
        return [], [f"mapping thiếu {channel}/{kind}"], None
    try:
        size = path.stat().st_size
        if size > MAX_FILE_BYTES:
            return [], [f"file > {MAX_FILE_BYTES // 2**20}MB — tách file"], None
        data = path.read_bytes()
    except OSError as e:
        return [], [f"không đọc được file: {type(e).__name__}"], None
    sha = hashlib.sha256(data).hexdigest()
    if not data or data.isspace():   # isspace: không copy như strip()
        return [], ["file rỗng"], sha
    try:
        return (*_parse(_table(data, suffix), channel, kind, spec), sha)
    except _TableError as e:
        return [], [str(e)], sha


def _parse(table, channel: str, kind: str, spec: dict
           ) -> tuple[list[dict], list[str]]:
    first = next(table, None)
    if first is None:
        return [], ["file rỗng"]
    hdr_no, header = first
    header = [None if h is None else _nfc(h) for h in header]
    col_of = {}   # field -> chỉ số cột
    want = {**{f: _nfc(h) for f, h in spec.get("optional", {}).items()},
            **{f: _nfc(h) for f, h in spec["required"].items()}}
    drop = {_nfc(h) for h in spec.get("drop", [])}
    by_header = {h: f for f, h in want.items()}
    errs = []
    seen = set()
    for j, h in enumerate(header):
        if h is None:
            errs.append(f"header cột {j + 1} trống")
            continue
        if h in seen:
            errs.append(f"header trùng: '{_safe(h)}'")
            continue
        seen.add(h)
        if h in by_header:
            col_of[by_header[h]] = j
        elif h not in drop:
            errs.append(f"cột lạ '{_safe(h)}' (không thuộc required/optional/drop "
                        "— thêm vào mapping sau khi xác minh)")
    missing = [h for f, h in spec["required"].items() if f not in col_of]
    if len(missing) == len(spec["required"]):
        # Không cột bắt buộc nào khớp → dòng đầu nhiều khả năng là DỮ LIỆU
        # (file thiếu header / sai loại export) — KHÔNG in ô (có thể là
        # tên/SĐT người mua), chỉ báo chung.
        msg = (f"dòng {hdr_no} không khớp header {channel}/{kind} nào "
               f"({len(header)} cột, 0/{len(missing)} cột bắt buộc) — "
               "sai kênh/loại export hoặc file thiếu dòng header")
        return [], [msg]
    if missing:
        errs.append("thiếu cột bắt buộc: " + ", ".join(
            f"'{_nfc(h)}'" for h in missing))
    if errs:
        return [], errs

    required = set(spec["required"])
    rows = []
    width = len(header)
    n_body = 0
    for row_no, cells in table:
        n_body += 1
        if any(c is not None for c in cells[width:]):
            errs.append(f"dòng {row_no}: thừa ô ngoài header")
            continue
        rec = {f: (cells[j] if j < len(cells) else None)
               for f, j in col_of.items()}
        out = {"row_no": row_no}
        bad = False
        for f in FIELDS:
            v = rec.get(f)
            if f in NUM_MAX:
                if v is None and f not in required:
                    out[f] = None
                    continue
                n = _to_int(v, NUM_MAX[f])
                if n is None:
                    errs.append(f"dòng {row_no}: field '{f}' không phải số "
                                "nguyên ≥ 0 hợp lệ")
                    bad = True
                out[f] = n
            elif f == "order_id":
                if kind != "order":
                    continue
                if v is None:
                    errs.append(f"dòng {row_no}: thiếu mã đơn")
                    bad = True
                    continue
                out["order_id_hash"] = hashlib.sha256(
                    f"{channel}:{v}".encode()).hexdigest()[:16]
            elif (f == "voucher_code" and v is not None
                  and len(v) > VOUCHER_MAX):
                errs.append(f"dòng {row_no}: field '{f}' dài quá "
                            f"{VOUCHER_MAX} ký tự")
                bad = True
            else:
                out[f] = v
        if out.get("item_id") is None and out.get("seller_sku") is None:
            errs.append(f"dòng {row_no}: thiếu cả item_id lẫn seller_sku")
            bad = True
        if not bad:
            rows.append(out)
    if not n_body:
        return [], ["file chỉ có header, không có dòng dữ liệu"]
    if errs:
        return [], errs
    return rows, []


# ---------- ánh xạ + đối chiếu (thuần) ----------

def map_sku(row: dict, channel: str, skus: dict) -> tuple[str | None, list]:
    """→ (sku_key | None, ứng viên). skus = {sku_key: {shopee_item_id,...}}."""
    ss = row.get("seller_sku")
    if ss and ss in skus:
        return ss, []
    iid = row.get("item_id")
    if channel == "shopee" and iid:
        cands = sorted(k for k, s in skus.items()
                       if s.get("shopee_item_id") == iid)
        if len(cands) == 1:
            return cands[0], []
        return None, cands
    return None, []


def _key(row: dict) -> tuple[str, str]:
    return (row.get("item_id") or "", row.get("seller_sku") or "")


def find_alerts(rows: list[dict], channel: str, kind: str, skus: dict,
                prices: dict, inactive: list[str]) -> dict:
    """Hàm thuần: → {(alert_kind, item_id, seller_sku): (sku_key, detail)}.
    prices = {sku_key: (price_vnd, price_id, approved_by)} giá list golden
    đang mở của kênh. Detail = giá trị theo nguồn, không có "đề xuất"."""
    out: dict = {}
    if kind not in ("listing", "ads"):
        return out
    inactive_n = {_nfc(s) for s in inactive}
    spend: dict = {}
    for r in rows:
        sku, cands = map_sku(r, channel, skus)
        k = _key(r)
        if sku is None:
            out[("item_unmapped", *k)] = (None, {
                "item_id": r.get("item_id"), "seller_sku": r.get("seller_sku"),
                "title": r.get("title"), "candidates": cands,
                "row_no": r["row_no"]})
            continue
        cat = skus[sku]["category"]
        if kind == "listing":
            st = r.get("status")
            if cat == "alcohol" and (st is None or _nfc(st) not in inactive_n):
                out[("alcohol_listed", *k)] = (sku, {
                    "sku_key": sku, "status_san": st,
                    "price_san": r.get("price"), "title": r.get("title"),
                    "row_no": r["row_no"]})
            g = prices.get(sku)
            if g is not None and r.get("price") != g[0]:
                out[("price_mismatch", *k)] = (sku, {
                    "sku_key": sku, "price_san": r.get("price"),
                    "price_golden": g[0], "price_id": g[1],
                    "golden_approved": g[2] is not None,
                    "approved_by": g[2],
                    "golden_note": ("giá golden đã duyệt" if g[2] else
                                    "giá golden CHƯA duyệt"),
                    "row_no": r["row_no"]})
        elif (cat == "tpbvsk" and skus[sku]["xnqc_status"] != "approved"
              and (r.get("ads_spend") or 0) > 0):
            spend[k] = spend.get(k, 0) + r["ads_spend"]
            out[("tpbvsk_ads_no_xnqc", *k)] = (sku, {
                "sku_key": sku, "ads_spend_vnd": spend[k],
                "xnqc_status": skus[sku]["xnqc_status"]})
    return out


# ---------- DB ----------

def _connect(dsn: str):
    return psycopg.connect(dsn, connect_timeout=10)


def _golden(conn, channel: str) -> tuple[dict, dict]:
    skus = {k: {"category": c, "xnqc_status": x, "shopee_item_id": i}
            for k, c, x, i in conn.execute(
                "select sku_key, category, xnqc_status, shopee_item_id"
                " from sku_master")}
    prices = {k: (p, i, a) for k, p, i, a in conn.execute(
        "select sku_key, price_vnd, id, approved_by from price_list"
        " where channel = %s and kind = 'list' and program_id is null"
        f" and valid_to is null and valid_from <= {TODAY_VN}", (channel,))}
    return skus, prices


def load(conn, rows: list[dict], channel: str, kind: str, file_name: str,
         sha: str) -> tuple[int, int]:
    """Insert snapshot + rows (không commit). → (snapshot_id, số row map)."""
    skus, _ = _golden(conn, channel)
    sid = conn.execute(
        "insert into market_snapshots (channel, export_kind, file_name,"
        " file_sha256, row_count) values (%s, %s, %s, %s, %s) returning id",
        (channel, kind, file_name, sha, len(rows))).fetchone()[0]
    params = []
    mapped = 0
    for r in rows:
        sku, _c = map_sku(r, channel, skus)
        mapped += sku is not None
        raw = {k: v for k, v in r.items() if k != "row_no"}
        params.append((sid, r["row_no"], channel, kind, r.get("item_id"),
                       r.get("seller_sku"), sku, r.get("title"),
                       r.get("price"), r.get("stock"), r.get("status"),
                       r.get("ads_spend"), r.get("qty"),
                       r.get("order_id_hash"), Jsonb(raw)))
    with conn.cursor() as cur:
        cur.executemany(
            "insert into market_rows (snapshot_id, row_no, channel,"
            " export_kind, item_id, seller_sku, sku_key, title, price_vnd,"
            " stock, status, ads_spend_vnd, qty, order_id_hash, raw) values"
            " (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
            params)
    return sid, mapped


def reconcile(conn, snapshot_id: int, inactive: list[str]) -> dict:
    """Mở/đóng `market_alerts` theo snapshot (không commit). Chỉ ĐỌC golden.
    → {kind: số cảnh báo open từ snapshot, "closed": số vừa đóng}."""
    channel, kind = conn.execute(
        "select channel, export_kind from market_snapshots where id = %s",
        (snapshot_id,)).fetchone()
    if kind not in ("listing", "ads"):
        return {}
    rows = [{"row_no": n, "item_id": i, "seller_sku": s, "title": t,
             "price": p, "status": st, "ads_spend": a}
            for n, i, s, t, p, st, a in conn.execute(
                "select row_no, item_id, seller_sku, title, price_vnd,"
                " status, ads_spend_vnd from market_rows"
                " where snapshot_id = %s order by row_no", (snapshot_id,))]
    skus, prices = _golden(conn, channel)
    found = find_alerts(rows, channel, kind, skus, prices, inactive)
    counts: dict = {}
    for (akind, iid, ss), (sku, detail) in found.items():
        conn.execute(
            "insert into market_alerts (kind, severity, channel, export_kind,"
            " item_id, seller_sku, sku_key, detail, snapshot_id) values"
            " (%s, %s, %s, %s, %s, %s, %s, %s, %s)"
            " on conflict (kind, channel, export_kind,"
            " (coalesce(item_id, '')), (coalesce(seller_sku, '')))"
            " do update set severity = excluded.severity,"
            " sku_key = excluded.sku_key, detail = excluded.detail,"
            " snapshot_id = excluded.snapshot_id, updated_at = now(),"
            " opened_at = case when market_alerts.status = 'closed'"
            " then now() else market_alerts.opened_at end,"
            " status = 'open', closed_at = null",
            (akind, SEVERITY[akind], channel, kind, iid or None, ss or None,
             sku, Jsonb(detail), snapshot_id))
        counts[akind] = counts.get(akind, 0) + 1
    # Chỉ đóng khi item CÓ MẶT trong file mới mà hết lỗi (bằng chứng đã sửa:
    # đã ẩn / giá khớp / chi phí 0). Item VẮNG mặt (xuất 1 phần, lọc danh
    # mục, thiếu trang) → giữ open — vắng ≠ đã sửa (lớp "file thiếu record
    # xóa state"; fail-closed với cảnh báo S1 pháp lý).
    keys = list(found)
    present = sorted({_key(r) for r in rows})
    counts["closed"] = conn.execute(
        "update market_alerts a set status = 'closed', closed_at = now(),"
        " updated_at = now() where a.status = 'open' and a.channel = %s"
        " and a.export_kind = %s and exists (select 1 from"
        " unnest(%s::text[], %s::text[]) as p(iid, ss)"
        " where p.iid = coalesce(a.item_id, '')"
        " and p.ss = coalesce(a.seller_sku, '')) and not exists (select 1"
        " from unnest(%s::text[], %s::text[], %s::text[]) as k(kind, iid, ss)"
        " where k.kind = a.kind and k.iid = coalesce(a.item_id, '')"
        " and k.ss = coalesce(a.seller_sku, ''))",
        (channel, kind, [k[0] for k in present], [k[1] for k in present],
         [k[0] for k in keys], [k[1] for k in keys],
         [k[2] for k in keys])).rowcount
    # open còn lại không thuộc snapshot này = item vắng mặt trong file.
    counts["kept_absent"] = conn.execute(
        "select count(*) from market_alerts where status = 'open'"
        " and channel = %s and export_kind = %s and snapshot_id <> %s",
        (channel, kind, snapshot_id)).fetchone()[0]
    return counts


def report(conn, mapping: dict | None = None) -> None:
    alerts = conn.execute(
        "select severity, kind, channel, export_kind, item_id, seller_sku,"
        " sku_key, detail, opened_at from market_alerts where status = 'open'"
        " order by severity, kind, channel, item_id, seller_sku").fetchall()
    print(f"[report] {len(alerts)} cảnh báo sàn đang mở (agent không sửa sàn"
          " — MXH sửa sau khi GĐ bán hàng/Pháp chế duyệt):")
    for sev, kind, ch, ek, iid, ss, sku, detail, opened in alerts:
        print(f"  [{sev}] {kind} {ch}/{ek} item={_clean(iid)} "
              f"seller_sku={_clean(ss)} "
              f"sku={sku} mở={opened:%Y-%m-%d %H:%M} "
              + _clean(json.dumps(detail, ensure_ascii=False,
                                  sort_keys=True), 2000))
    snaps = conn.execute(
        "select distinct on (channel, export_kind) id, channel, export_kind,"
        " file_name, row_count, imported_at,"
        f" {TODAY_VN} - (imported_at at time zone 'Asia/Ho_Chi_Minh')::date"
        " from market_snapshots order by channel, export_kind,"
        " imported_at desc, id desc").fetchall()
    print(f"[report] snapshot mới nhất: {len(snaps)}")
    latest = {}
    for sid, ch, ek, fname, n, ts, age in snaps:
        latest[(ch, ek)] = sid
        print(f"  {ch}/{ek} #{sid} {_clean(fname)} {n} dòng, nạp "
              f"{ts:%Y-%m-%d %H:%M} ({age} ngày trước)")
    for ch in CHANNELS:
        if (ch, "listing") in latest and (ch, "ads") not in latest:
            print(f"[warn] {ch}: không đủ dữ liệu quảng cáo (có listing, "
                  "chưa có file ads) — chưa kiểm được TPBVSK chạy ads")
        if (ch, "listing") in latest:
            n_nop = conn.execute(
                "select count(*) from market_rows r where r.snapshot_id = %s"
                " and r.sku_key is not null and not exists (select 1 from"
                " price_list p where p.sku_key = r.sku_key and p.channel = %s"
                " and p.kind = 'list' and p.program_id is null and"
                f" p.valid_to is null and p.valid_from <= {TODAY_VN})",
                (latest[(ch, "listing")], ch)).fetchone()[0]
            print(f"  {ch}: {n_nop} dòng listing đã map nhưng golden chưa có "
                  "giá list để so")
        if (ch, "order") in latest:
            cancel = {_nfc(s) for s in ((mapping or {}).get(ch, {})
                                        .get("order", {})
                                        .get("cancel_status", []))}
            by = conn.execute(
                "select coalesce(status, '(trống)'), count(distinct"
                " order_id_hash), count(*) from market_rows"
                " where snapshot_id = %s group by 1 order by 1",
                (latest[(ch, "order")],)).fetchall()
            n_cancel = sum(o for s, o, _ in by if _nfc(s) in cancel)
            print(f"  {ch} đơn: hủy={n_cancel} đơn; theo trạng thái: "
                  + ", ".join(f"{_clean(s, 40)}={o} đơn/{n} dòng" for s, o, n in by))


# ---------- CLI ----------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        prog="python -m ingest.market_csv",
        description="Nạp file xuất Seller Center (chỉ đọc) / báo cáo cảnh "
                    "báo sàn.")
    ap.add_argument("--file", type=Path, help="file .csv / .xlsx đã xuất")
    ap.add_argument("--channel", choices=CHANNELS)
    ap.add_argument("--kind", choices=KINDS)
    ap.add_argument("--report", action="store_true",
                    help="in cảnh báo open + snapshot mới nhất")
    args = ap.parse_args(argv)

    n_ops = sum([args.file is not None, args.report])
    if n_ops > 1:
        ap.error("chỉ được 1 thao tác mỗi lần (--file | --report)")
    if n_ops == 0:
        ap.error("cần --file ... hoặc --report")
    if args.file is not None and (args.channel is None or args.kind is None):
        ap.error("--file cần đủ --channel và --kind")
    if args.report and (args.channel or args.kind):
        ap.error("--channel/--kind chỉ dùng với --file")

    mapping, why = load_mapping()
    if args.file is not None:
        if mapping is None:
            print(f"[err] {why}", flush=True)
            return 1
        rows, errs, sha = read_export(args.file, args.channel, args.kind,
                                      mapping)
        if errs:
            for e in errs[:MAX_PRINT_ERRS]:
                print(f"[err] {e}", flush=True)
            if len(errs) > MAX_PRINT_ERRS:
                print(f"[err] ... và {len(errs) - MAX_PRINT_ERRS} lỗi khác")
            print(f"[err] {len(errs)} lỗi — không connect, không nạp gì.",
                  flush=True)
            return 1

    load_dotenv(ROOT / ".env")
    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        print("[err] thiếu DATABASE_URL trong .env")
        return 1
    try:
        conn = _connect(dsn)
    except psycopg.OperationalError as e:
        print(f"[err] không kết nối được Postgres: {type(e).__name__}")
        return 1
    try:
        if args.report:
            report(conn, mapping)
            return 0
        old = conn.execute(
            "select id, channel, export_kind from market_snapshots"
            " where file_sha256 = %s", (sha,)).fetchone()
        if old:
            print(f"[skip] file đã nạp: snapshot #{old[0]} {old[1]}/{old[2]}"
                  " — không ghi gì.")
            return 0
        sid, mapped = load(conn, rows, args.channel, args.kind,
                           args.file.name, sha)
        inactive = mapping[args.channel][args.kind].get("inactive_status", [])
        alerts = reconcile(conn, sid, inactive)
        conn.commit()
    except psycopg.errors.UniqueViolation as e:
        conn.rollback()
        if "file_sha256" in (e.diag.constraint_name or ""):
            print("[skip] file vừa được nạp song song — không ghi gì.")
            return 0
        print(f"[err] lỗi DB, đã rollback: {type(e).__name__}")
        return 1
    except psycopg.errors.UndefinedTable:
        conn.rollback()
        print("[err] thiếu bảng market_* / golden — chạy `python "
              "scripts/apply_schema.py` trước.")
        return 1
    except psycopg.Error as e:
        conn.rollback()
        print(f"[err] lỗi DB, đã rollback: {type(e).__name__}")
        return 1
    finally:
        conn.close()
    print(f"[load] snapshot #{sid} {args.channel}/{args.kind} "
          f"{_clean(args.file.name)}: {len(rows)} dòng, map golden {mapped}")
    kept = alerts.pop("kept_absent", 0)
    print("[alerts] " + (", ".join(f"{k}={v}" for k, v in
                                   sorted(alerts.items()))
                         or "export order — không đối chiếu ở P1"))
    if kept:
        print(f"[warn] {kept} cảnh báo {args.channel}/{args.kind} không thấy "
              "item trong file — GIỮ MỞ (vắng ≠ đã sửa; xuất TOÀN BỘ sản "
              "phẩm, không lọc, để đóng)")
    print("[done]")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
